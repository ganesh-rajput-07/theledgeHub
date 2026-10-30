import io

import docx

from app.ingestion.extractors.base import ExtractedPage, ExtractionError


def extract_docx(data: bytes) -> list[ExtractedPage]:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX: {exc}") from exc

    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return [ExtractedPage(page_number=None, text="\n\n".join(blocks))]
